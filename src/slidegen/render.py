"""Turn validated specs into a `.pptx`.

The renderer adds a slide per spec, applies the chosen layout, and pours text
into placeholders. It sets no position, size, font or colour — every one of
those is inherited from the layout, which is the whole point.
"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path

from pptx import Presentation as open_presentation
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

from .spec import PlaceholderValue, SlideSpec, ensure_valid
from .template import LayoutCatalog, load_template

logger = logging.getLogger(__name__)

TemplateSource = PresentationType | str | Path | None


def _fresh_presentation(template: TemplateSource) -> PresentationType:
    """A private copy of the template, so rendering never mutates the caller's.

    The Streamlit app caches one loaded template across reruns; without the copy,
    every render would append to the previous one.
    """
    if isinstance(template, PresentationType):
        buffer = BytesIO()
        template.save(buffer)
        buffer.seek(0)
        return open_presentation(buffer)
    return load_template(template)


def _drop_existing_slides(presentation: PresentationType) -> None:
    """Start from an empty deck.

    A `.potx` normally carries no slides, but a `.pptx` pressed into service as a
    template often does, and those are not part of the output.
    """
    slide_id_list = presentation.slides._sldIdLst  # noqa: SLF001 — no public API
    for slide_id in list(slide_id_list):
        presentation.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def _fill_text_frame(placeholder, value: PlaceholderValue) -> None:
    """Write text into a placeholder, leaving all formatting to the layout."""
    text_frame = placeholder.text_frame
    if isinstance(value, str):
        text_frame.text = value
        return

    items = [item for item in value if item.strip()]
    text_frame.clear()
    if not items:
        return
    text_frame.paragraphs[0].text = items[0]
    text_frame.paragraphs[0].level = 0
    for item in items[1:]:
        paragraph = text_frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0


def _apply(spec: SlideSpec, slide: Slide) -> None:
    by_idx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    for idx, value in spec.placeholders.items():
        placeholder = by_idx.get(idx)
        if placeholder is None:
            # Validation already ruled this out; a layout that changed underneath
            # us is not worth failing a render over.
            logger.warning("placeholder idx %s absent from slide, skipping", idx)
            continue
        _fill_text_frame(placeholder, value)

    if spec.notes:
        slide.notes_slide.notes_text_frame.text = spec.notes


def build(
    specs: list[SlideSpec],
    catalog: LayoutCatalog,
    template: TemplateSource = None,
    *,
    validate: bool = True,
) -> PresentationType:
    """Build the presentation object. Use `render` for bytes."""
    if validate:
        ensure_valid(specs, catalog)

    presentation = _fresh_presentation(template)
    _drop_existing_slides(presentation)

    for spec in specs:
        layout = presentation.slide_layouts[spec.layout_index]
        slide = presentation.slides.add_slide(layout)
        _apply(spec, slide)

    logger.debug(
        "rendered deck",
        extra={"slides": len(specs), "template": catalog.source},
    )
    return presentation


def render(
    specs: list[SlideSpec],
    catalog: LayoutCatalog,
    template: TemplateSource = None,
    *,
    validate: bool = True,
) -> bytes:
    """Render specs to `.pptx` bytes, one slide per spec."""
    presentation = build(specs, catalog, template, validate=validate)
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
