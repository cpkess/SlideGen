"""The `--brief` entry point, run end to end against the mock provider."""

from __future__ import annotations

import json
from io import BytesIO, StringIO

import pytest
from pptx import Presentation

from slidegen.cli import main

BRIEF = (
    "Connected services needs a Q3 exec summary. "
    "Adoption grew 24 percent. Churn is flat at 3 percent."
)


def test_it_writes_a_deck(tmp_path, capsys):
    out = tmp_path / "deck.pptx"

    assert main(["--brief", BRIEF, "--out", str(out)]) == 0

    deck = Presentation(BytesIO(out.read_bytes()))
    assert len(deck.slides) == 1
    assert deck.slides[0].shapes.title.text_frame.text.strip()
    assert "Wrote" in capsys.readouterr().out


def test_it_lists_candidates_without_writing(capsys):
    assert main(["--brief", BRIEF, "--list"]) == 0

    out = capsys.readouterr().out
    assert out.count("layout_index") == 3


def test_candidates_can_be_emitted_as_json(capsys):
    assert main(["--brief", BRIEF, "--list", "--json"]) == 0

    payload = json.loads(capsys.readouterr().out)
    assert len(payload) == 3
    assert {"layout_index", "rationale", "placeholders", "notes"} == set(payload[0])


def test_a_specific_candidate_can_be_chosen(tmp_path):
    first = tmp_path / "one.pptx"
    third = tmp_path / "three.pptx"

    main(["--brief", BRIEF, "--out", str(first), "--candidate", "1", "-q"])
    main(["--brief", BRIEF, "--out", str(third), "--candidate", "3", "-q"])

    layout_of = lambda path: Presentation(BytesIO(path.read_bytes())).slides[0].slide_layout.name
    assert layout_of(first) != layout_of(third)


def test_all_candidates_can_go_into_one_deck(tmp_path):
    out = tmp_path / "all.pptx"

    assert main(["--brief", BRIEF, "--out", str(out), "--all", "-q"]) == 0

    assert len(Presentation(BytesIO(out.read_bytes())).slides) == 3


def test_the_brief_can_come_from_a_file(tmp_path):
    brief_file = tmp_path / "brief.txt"
    brief_file.write_text(BRIEF, encoding="utf-8")
    out = tmp_path / "deck.pptx"

    assert main(["--brief-file", str(brief_file), "--out", str(out), "-q"]) == 0
    assert out.exists()


def test_the_brief_can_come_from_stdin(tmp_path, monkeypatch):
    monkeypatch.setattr("sys.stdin", StringIO(BRIEF))
    out = tmp_path / "deck.pptx"

    assert main(["--brief", "-", "--out", str(out), "-q"]) == 0
    assert out.exists()


def test_an_empty_brief_is_rejected(capsys):
    assert main(["--brief", "   "]) == 2
    assert "empty" in capsys.readouterr().err


def test_a_missing_out_path_is_rejected(capsys):
    assert main(["--brief", BRIEF, "-q"]) == 2
    assert "--out is required" in capsys.readouterr().err


def test_an_out_of_range_candidate_is_rejected(tmp_path, capsys):
    code = main(["--brief", BRIEF, "--out", str(tmp_path / "d.pptx"), "--candidate", "9", "-q"])

    assert code == 2
    assert "out of range" in capsys.readouterr().err


def test_a_missing_template_is_reported(tmp_path, capsys):
    code = main(["--brief", BRIEF, "--template", str(tmp_path / "nope.potx"), "--list"])

    assert code == 2
    assert "error" in capsys.readouterr().err


def test_overflow_warnings_go_to_stderr(tmp_path, capsys):
    out = tmp_path / "deck.pptx"

    main(["--brief", "Padding padding padding. " * 200, "--out", str(out), "-q"])

    assert "warning:" in capsys.readouterr().err
    assert out.exists()  # a warning, not a failure


def test_brief_and_brief_file_are_mutually_exclusive():
    with pytest.raises(SystemExit):
        main(["--brief", BRIEF, "--brief-file", "x.txt"])
