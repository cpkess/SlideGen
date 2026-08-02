"""The mock provider is what makes `pytest` and the demo path work offline."""

from __future__ import annotations

from io import BytesIO

import pytest
from pptx import Presentation

from slidegen.llm.base import ProviderError
from slidegen.llm.mock import MockProvider
from slidegen.render import render
from slidegen.spec import validate_specs
from slidegen.template import LayoutCatalog

BRIEF = """Connected services needs a Q3 exec summary.
Adoption grew 24 percent quarter over quarter.
Churn is flat at 3 percent.
Dealer onboarding capacity is the main risk.
We need two more integration engineers by October."""


def test_it_returns_the_requested_number_of_candidates(catalog):
    assert len(MockProvider().generate(BRIEF, catalog, 3)) == 3
    assert len(MockProvider().generate(BRIEF, catalog, 1)) == 1


def test_every_candidate_validates(catalog):
    specs = MockProvider().generate(BRIEF, catalog, 3)

    assert validate_specs(specs, catalog) == {}


def test_candidates_use_different_layouts(catalog):
    specs = MockProvider().generate(BRIEF, catalog, 3)

    assert len({spec.layout_index for spec in specs}) == 3


def test_candidates_differ_in_shape_not_only_wording(catalog):
    specs = MockProvider().generate(BRIEF, catalog, 3)

    shapes = {
        (spec.layout_index, tuple(type(value).__name__ for value in spec.placeholders.values()))
        for spec in specs
    }
    assert len(shapes) == 3


def test_it_is_deterministic(catalog):
    first = MockProvider().generate(BRIEF, catalog, 3)
    second = MockProvider().generate(BRIEF, catalog, 3)

    assert first == second


def test_the_brief_reaches_the_slides(catalog):
    specs = MockProvider().generate(BRIEF, catalog, 3)

    assert all("Connected services" in spec.text_for(0) for spec in specs)
    assert any("24 percent" in spec.text_for(idx) for spec in specs for idx in spec.placeholders)


def test_it_round_trips_through_validation_and_render(catalog, presentation):
    specs = MockProvider().generate(BRIEF, catalog, 3)

    deck = Presentation(BytesIO(render(specs, catalog, presentation)))

    assert len(deck.slides) == 3
    for slide in deck.slides:
        assert slide.shapes.title.text_frame.text.strip()


def test_an_empty_brief_still_produces_valid_candidates(catalog):
    specs = MockProvider().generate("", catalog, 3)

    assert validate_specs(specs, catalog) == {}
    assert specs[0].text_for(0) == "Untitled brief"


def test_a_one_sentence_brief_still_fills_the_body(catalog):
    specs = MockProvider().generate("Ship it.", catalog, 3)

    assert validate_specs(specs, catalog) == {}


def test_long_sentences_are_truncated_not_dropped(catalog):
    specs = MockProvider().generate("word " * 400, catalog, 1)

    assert specs[0].text_for(0).endswith("…")


def test_a_template_with_nothing_to_fill_is_reported_clearly(catalog):
    empty = LayoutCatalog(source="stub", layouts=[])

    with pytest.raises(ProviderError, match="no layout with a fillable text placeholder"):
        MockProvider().generate(BRIEF, empty, 3)


def test_it_asks_for_more_candidates_than_there_are_layouts(catalog):
    """The picker must not crash on a thin template."""
    trimmed = catalog.model_copy(update={"layouts": catalog.layouts[:2]})

    specs = MockProvider().generate(BRIEF, trimmed, 3)

    assert len(specs) == 3
    assert validate_specs(specs, trimmed) == {}
